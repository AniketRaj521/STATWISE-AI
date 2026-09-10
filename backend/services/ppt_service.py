import os
import json
import re
from pathlib import Path
from typing import Any, Dict, List

from dotenv import load_dotenv
from groq import Groq

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


# ============================================================
# ENVIRONMENT
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent
ENV_PATH = BASE_DIR / ".env"

load_dotenv(ENV_PATH, override=True)

GROQ_API_KEY = os.getenv("GROQ_API_KEY")

if not GROQ_API_KEY:
    print("WARNING: GROQ_API_KEY is not configured.")


# ============================================================
# DIRECTORIES
# ============================================================

GENERATED_DIR = BASE_DIR / "generated_ppts"
GENERATED_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# BRAND COLORS
# ============================================================

NAVY = RGBColor(23, 32, 51)
DARK_NAVY = RGBColor(14, 22, 38)
CREAM = RGBColor(255, 253, 245)
YELLOW = RGBColor(246, 215, 106)
LIGHT_YELLOW = RGBColor(255, 247, 198)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(34, 197, 94)
RED = RGBColor(239, 68, 68)
BLUE = RGBColor(59, 130, 246)
PURPLE = RGBColor(139, 92, 246)
GRAY = RGBColor(100, 116, 139)
LIGHT_GRAY = RGBColor(241, 245, 249)
MID_GRAY = RGBColor(203, 213, 225)


# ============================================================
# HELPERS
# ============================================================

def clean_filename(filename: str) -> str:
    """
    Make a safe filename for Windows/Linux.
    """
    filename = os.path.basename(filename)

    filename = re.sub(
        r"[^a-zA-Z0-9._-]",
        "_",
        filename
    )

    filename = re.sub(
        r"_+",
        "_",
        filename
    )

    return filename[:120]


def clean_json_response(text: str) -> Dict[str, Any]:
    """
    Extract JSON from a Groq response.
    Handles markdown fences and accidental surrounding text.
    """

    if not text:
        raise ValueError("AI returned an empty response.")

    text = text.strip()

    # Remove markdown code fence
    text = re.sub(
        r"^```(?:json)?\s*",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"\s*```$",
        "",
        text
    )

    # Direct parse
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    # Find first { and last }
    start = text.find("{")
    end = text.rfind("}")

    if start == -1 or end == -1 or end <= start:
        raise ValueError(
            "Could not find valid JSON in AI response."
        )

    candidate = text[start:end + 1]

    return json.loads(candidate)


def safe_text(value: Any, default: str = "") -> str:
    if value is None:
        return default

    if isinstance(value, str):
        return value.strip()

    return str(value).strip()


def safe_list(value: Any) -> List[Any]:
    if isinstance(value, list):
        return value

    return []


# ============================================================
# NORMALIZE AI PRESENTATION
# ============================================================

ALLOWED_TYPES = {
    "title",
    "objectives",
    "concept",
    "comparison",
    "process",
    "timeline",
    "components",
    "table",
    "chart",
    "case_study",
    "key_insight",
    "quiz",
    "summary",
    "references",
}


def normalize_slide(slide: Dict[str, Any], number: int) -> Dict[str, Any]:

    if not isinstance(slide, dict):
        slide = {}

    slide_type = safe_text(
        slide.get("type"),
        "concept"
    ).lower()

    if slide_type not in ALLOWED_TYPES:
        slide_type = "concept"

    title = safe_text(
        slide.get("title"),
        f"Key Topic {number}"
    )

    subtitle = safe_text(
        slide.get("subtitle")
    )

    points = safe_list(
        slide.get("points")
    )

    points = [
        safe_text(item)
        for item in points
        if safe_text(item)
    ][:8]

    items = safe_list(
        slide.get("items")
    )

    normalized_items = []

    for item in items:

        if isinstance(item, dict):

            normalized_items.append({
                "name": safe_text(
                    item.get("name"),
                    "Item"
                ),
                "description": safe_text(
                    item.get("description")
                ),
                "example": safe_text(
                    item.get("example")
                ),
                "value": safe_text(
                    item.get("value")
                ),
            })

        else:

            normalized_items.append({
                "name": safe_text(item),
                "description": "",
                "example": "",
                "value": "",
            })

    data = safe_list(
        slide.get("data")
    )

    normalized_data = []

    for item in data:

        if isinstance(item, dict):

            label = safe_text(
                item.get("label"),
                "Item"
            )

            value = item.get("value", 0)

            try:
                value = float(value)
            except Exception:
                value = 0

            normalized_data.append({
                "label": label,
                "value": value
            })

    questions = safe_list(
        slide.get("questions")
    )

    return {
        "slide_number": number,
        "type": slide_type,
        "title": title,
        "subtitle": subtitle,
        "points": points,
        "items": normalized_items[:8],
        "data": normalized_data[:8],
        "questions": [
            safe_text(q)
            for q in questions
            if safe_text(q)
        ][:5],
        "answer": safe_text(
            slide.get("answer")
        ),
        "insight": safe_text(
            slide.get("insight")
        ),
        "source": safe_text(
            slide.get("source")
        ),
    }


def normalize_presentation(data: Dict[str, Any]) -> Dict[str, Any]:

    if not isinstance(data, dict):
        data = {}

    raw_slides = safe_list(
        data.get("slides")
    )

    slides = []

    for index, slide in enumerate(
        raw_slides,
        start=1
    ):
        slides.append(
            normalize_slide(
                slide,
                index
            )
        )

    return {
        "title": safe_text(
            data.get("title"),
            "STATWISE AI Learning Presentation"
        ),
        "subtitle": safe_text(
            data.get("subtitle"),
            "AI-generated presentation from learning material"
        ),
        "topic": safe_text(
            data.get("topic"),
            "Learning Material"
        ),
        "audience": safe_text(
            data.get("audience"),
            "Students and learners"
        ),
        "style": safe_text(
            data.get("style"),
            "Visual Learning"
        ),
        "learning_objectives": [
            safe_text(x)
            for x in safe_list(
                data.get("learning_objectives")
            )
            if safe_text(x)
        ][:6],
        "key_statistics": [
            safe_text(x)
            for x in safe_list(
                data.get("key_statistics")
            )
            if safe_text(x)
        ][:8],
        "slides": slides,
    }


# ============================================================
# AI PRESENTATION PLANNER
# ============================================================

def generate_presentation_plan(
    pdf_text: str,
    style: str = "Visual Learning"
) -> Dict[str, Any]:

    if not pdf_text or not pdf_text.strip():
        return {
            "success": False,
            "message": "No readable PDF text was provided."
        }

    if not GROQ_API_KEY:
        return {
            "success": False,
            "message": "GROQ_API_KEY is not configured."
        }

    # Keep enough content while protecting token usage.
    source_text = pdf_text[:65000]

    client = Groq(
        api_key=GROQ_API_KEY
    )

    prompt = f"""
You are the presentation intelligence engine of STATWISE AI.

Your task is to transform the supplied PDF learning material into
a rich, professional, visually structured educational presentation.

IMPORTANT RULES:

1. Use ONLY information supported by the supplied PDF.
2. Do not invent facts, statistics, dates, research findings,
   citations, or examples that are not reasonably supported by
   the supplied material.
3. Create approximately 12 to 15 slides.
4. The presentation should contain substantial information.
5. Avoid repeating the same information on multiple slides.
6. Select different slide types when appropriate.
7. Make the presentation useful for learning and revision.
8. Include important definitions, concepts, relationships,
   processes, comparisons, examples, applications, and
   takeaways when present in the source.
9. If the PDF contains numerical information suitable for a chart,
   create a chart slide. Otherwise do not fabricate numerical data.
10. Include a knowledge-check / quiz slide near the end.
11. Include a summary slide.
12. Include references only for sources actually identifiable
    from the PDF.
13. Style requested: {style}
14. Return ONLY valid JSON.
15. Do NOT use markdown.
16. Do NOT put comments outside the JSON.

Allowed slide types:

title
objectives
concept
comparison
process
timeline
components
table
chart
case_study
key_insight
quiz
summary
references

Return exactly this structure:

{{
  "title": "Presentation title",
  "subtitle": "Short subtitle",
  "topic": "Main topic",
  "audience": "Target audience",
  "style": "{style}",
  "learning_objectives": [
    "objective 1",
    "objective 2",
    "objective 3"
  ],
  "key_statistics": [],
  "slides": [
    {{
      "type": "title",
      "title": "Title",
      "subtitle": "Subtitle",
      "points": [],
      "items": [],
      "data": [],
      "questions": [],
      "answer": "",
      "insight": "",
      "source": ""
    }}
  ]
}}

For comparison slides use items such as:

{{
  "name": "LAN",
  "description": "Local Area Network",
  "example": "Office or campus",
  "value": ""
}}

For chart slides use:

{{
  "label": "Category",
  "value": 25
}}

For quiz slides use:

"questions": [
  "Question 1",
  "Question 2"
],
"answer": "Answer explanation"

SOURCE MATERIAL:

{source_text}
"""

    try:

        completion = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a precise educational presentation "
                        "planner. Output valid JSON only."
                    )
                },
                {
                    "role": "user",
                    "content": prompt
                }
            ],
            temperature=0.2,
            max_tokens=12000,
        )

        raw = completion.choices[0].message.content

        print("\n========== PPT AI RESPONSE ==========\n")
        print(raw[:5000])

        data = clean_json_response(raw)

        presentation = normalize_presentation(
            data
        )

        if len(presentation["slides"]) < 5:

            return {
                "success": False,
                "message": (
                    "AI generated too few slides. "
                    "Please try again."
                )
            }

        return {
            "success": True,
            "presentation": presentation
        }

    except Exception as error:

        print(
            "\n========== PPT AI ERROR ==========\n",
            error
        )

        return {
            "success": False,
            "message": "Failed to generate presentation plan.",
            "error": str(error)
        }


# ============================================================
# POWERPOINT HELPERS
# ============================================================

def set_slide_background(
    slide,
    color=CREAM
):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_top_bar(slide):

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        Inches(13.333),
        Inches(0.12)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = YELLOW
    shape.line.fill.background()


def add_footer(
    slide,
    slide_number: int
):

    # Bottom line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.55),
        Inches(7.12),
        Inches(12.2),
        Inches(0.015)
    )

    line.fill.solid()
    line.fill.fore_color.rgb = MID_GRAY
    line.line.fill.background()

    # Brand
    box = slide.shapes.add_textbox(
        Inches(0.6),
        Inches(7.18),
        Inches(4),
        Inches(0.25)
    )

    p = box.text_frame.paragraphs[0]

    p.text = "STATWISE AI"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = NAVY

    # Slide number
    number_box = slide.shapes.add_textbox(
        Inches(11.8),
        Inches(7.16),
        Inches(0.8),
        Inches(0.25)
    )

    p = number_box.text_frame.paragraphs[0]

    p.text = str(slide_number)
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(9)
    p.font.color.rgb = GRAY


def add_title(
    slide,
    title: str,
    subtitle: str = ""
):

    title_box = slide.shapes.add_textbox(
        Inches(0.65),
        Inches(0.42),
        Inches(11.9),
        Inches(0.65)
    )

    tf = title_box.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(27)
    p.font.bold = True
    p.font.color.rgb = NAVY

    if subtitle:

        subtitle_box = slide.shapes.add_textbox(
            Inches(0.67),
            Inches(1.02),
            Inches(11.5),
            Inches(0.4)
        )

        p = subtitle_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY


def add_bullet_card(
    slide,
    x,
    y,
    w,
    h,
    title,
    points,
    accent=YELLOW
):

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    card.fill.solid()
    card.fill.fore_color.rgb = WHITE

    card.line.color.rgb = MID_GRAY

    # Accent
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(0.08),
        Inches(h)
    )

    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()

    text_box = slide.shapes.add_textbox(
        Inches(x + 0.25),
        Inches(y + 0.2),
        Inches(w - 0.45),
        Inches(h - 0.35)
    )

    tf = text_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    for point in points:

        p = tf.add_paragraph()

        p.text = f"• {point}"
        p.font.size = Pt(12)
        p.font.color.rgb = NAVY
        p.space_before = Pt(8)


def add_badge(
    slide,
    x,
    y,
    text,
    color=YELLOW
):

    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(1.6),
        Inches(0.42)
    )

    badge.fill.solid()
    badge.fill.fore_color.rgb = color
    badge.line.fill.background()

    box = slide.shapes.add_textbox(
        Inches(x + 0.08),
        Inches(y + 0.07),
        Inches(1.44),
        Inches(0.25)
    )

    p = box.text_frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = NAVY


def add_process_nodes(
    slide,
    items
):

    items = items[:6]

    if not items:
        return

    start_x = 0.75
    gap = 2.05

    for index, item in enumerate(items):

        x = start_x + index * gap

        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x),
            Inches(2.4),
            Inches(0.85),
            Inches(0.85)
        )

        circle.fill.solid()
        circle.fill.fore_color.rgb = (
            YELLOW if index % 2 == 0
            else LIGHT_YELLOW
        )

        circle.line.color.rgb = NAVY

        number = slide.shapes.add_textbox(
            Inches(x),
            Inches(2.58),
            Inches(0.85),
            Inches(0.3)
        )

        p = number.text_frame.paragraphs[0]
        p.text = str(index + 1)
        p.alignment = PP_ALIGN.CENTER
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = NAVY

        name = slide.shapes.add_textbox(
            Inches(x - 0.45),
            Inches(3.38),
            Inches(1.75),
            Inches(0.7)
        )

        p = name.text_frame.paragraphs[0]
        p.text = safe_text(
            item.get("name"),
            f"Step {index + 1}"
        )
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = NAVY

        if index < len(items) - 1:

            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(x + 0.95),
                Inches(2.65),
                Inches(0.9),
                Inches(0.28)
            )

            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GRAY
            arrow.line.fill.background()


def add_comparison_cards(
    slide,
    items
):

    items = items[:4]

    if not items:
        return

    card_width = 2.85
    gap = 0.25

    for index, item in enumerate(items):

        x = 0.6 + index * (
            card_width + gap
        )

        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(1.75),
            Inches(card_width),
            Inches(4.75)
        )

        card.fill.solid()
        card.fill.fore_color.rgb = WHITE

        card.line.color.rgb = (
            YELLOW if index == 0
            else MID_GRAY
        )

        # Header
        header = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x + 0.18),
            Inches(1.98),
            Inches(card_width - 0.36),
            Inches(0.62)
        )

        header.fill.solid()
        header.fill.fore_color.rgb = (
            LIGHT_YELLOW if index == 0
            else LIGHT_GRAY
        )

        header.line.fill.background()

        name_box = slide.shapes.add_textbox(
            Inches(x + 0.25),
            Inches(2.13),
            Inches(card_width - 0.5),
            Inches(0.3)
        )

        p = name_box.text_frame.paragraphs[0]
        p.text = safe_text(
            item.get("name"),
            "Category"
        )
        p.alignment = PP_ALIGN.CENTER
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = NAVY

        body = slide.shapes.add_textbox(
            Inches(x + 0.28),
            Inches(2.85),
            Inches(card_width - 0.55),
            Inches(3.2)
        )

        tf = body.text_frame
        tf.word_wrap = True

        description = safe_text(
            item.get("description")
        )

        example = safe_text(
            item.get("example")
        )

        value = safe_text(
            item.get("value")
        )

        fields = []

        if description:
            fields.append(
                f"Description\n{description}"
            )

        if example:
            fields.append(
                f"Example\n{example}"
            )

        if value:
            fields.append(
                f"Value\n{value}"
            )

        for idx, text in enumerate(fields):

            p = (
                tf.paragraphs[0]
                if idx == 0
                else tf.add_paragraph()
            )

            p.text = text
            p.font.size = Pt(11)
            p.font.color.rgb = NAVY
            p.space_after = Pt(16)


def add_table_slide(
    slide,
    items
):

    items = items[:6]

    if not items:
        return

    columns = [
        "Concept",
        "Description",
        "Example"
    ]

    rows = len(items) + 1

    table_shape = slide.shapes.add_table(
        rows,
        3,
        Inches(0.65),
        Inches(1.75),
        Inches(12),
        Inches(4.9)
    )

    table = table_shape.table

    table.columns[0].width = Inches(2.3)
    table.columns[1].width = Inches(5.0)
    table.columns[2].width = Inches(4.7)

    for col, heading in enumerate(columns):

        cell = table.cell(0, col)

        cell.text = heading

        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY

        for p in cell.text_frame.paragraphs:

            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER

    for row_index, item in enumerate(
        items,
        start=1
    ):

        values = [
            safe_text(
                item.get("name"),
                "Concept"
            ),
            safe_text(
                item.get("description"),
                "—"
            ),
            safe_text(
                item.get("example"),
                "—"
            ),
        ]

        for col, value in enumerate(values):

            cell = table.cell(
                row_index,
                col
            )

            cell.text = value
            cell.fill.solid()

            cell.fill.fore_color.rgb = (
                WHITE
                if row_index % 2
                else LIGHT_GRAY
            )

            for p in cell.text_frame.paragraphs:

                p.font.size = Pt(10)
                p.font.color.rgb = NAVY


def add_chart_slide(
    slide,
    data
):

    data = data[:6]

    if not data:
        return False

    max_value = max(
        [item["value"] for item in data] + [1]
    )

    start_x = 1.0
    base_y = 5.9

    chart_height = 3.4

    bar_width = 1.35
    gap = 0.55

    for index, item in enumerate(data):

        x = start_x + index * (
            bar_width + gap
        )

        height = (
            item["value"] / max_value
        ) * chart_height

        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(base_y - height),
            Inches(bar_width),
            Inches(height)
        )

        bar.fill.solid()

        bar.fill.fore_color.rgb = (
            YELLOW
            if index % 2 == 0
            else BLUE
        )

        bar.line.fill.background()

        value_box = slide.shapes.add_textbox(
            Inches(x),
            Inches(base_y - height - 0.35),
            Inches(bar_width),
            Inches(0.3)
        )

        p = value_box.text_frame.paragraphs[0]
        p.text = str(
            round(item["value"], 1)
        )
        p.alignment = PP_ALIGN.CENTER
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = NAVY

        label_box = slide.shapes.add_textbox(
            Inches(x - 0.15),
            Inches(base_y + 0.15),
            Inches(bar_width + 0.3),
            Inches(0.6)
        )

        p = label_box.text_frame.paragraphs[0]
        p.text = item["label"]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(9)
        p.font.color.rgb = NAVY

    return True


# ============================================================
# SLIDE GENERATORS
# ============================================================

def create_title_slide(
    prs,
    presentation,
    slide_number
):

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_slide_background(
        slide,
        NAVY
    )

    # Large decorative circles
    circle1 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(9.8),
        Inches(-1),
        Inches(4.5),
        Inches(4.5)
    )

    circle1.fill.solid()
    circle1.fill.fore_color.rgb = YELLOW
    circle1.fill.transparency = 25
    circle1.line.fill.background()

    circle2 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(-1.5),
        Inches(5.1),
        Inches(4),
        Inches(4)
    )

    circle2.fill.solid()
    circle2.fill.fore_color.rgb = BLUE
    circle2.fill.transparency = 50
    circle2.line.fill.background()

    add_badge(
        slide,
        0.8,
        1.0,
        "STATWISE AI",
        YELLOW
    )

    title = slide.shapes.add_textbox(
        Inches(0.8),
        Inches(2.0),
        Inches(10.8),
        Inches(1.5)
    )

    tf = title.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = presentation["title"]
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = WHITE

    subtitle = slide.shapes.add_textbox(
        Inches(0.82),
        Inches(3.55),
        Inches(9.5),
        Inches(1.0)
    )

    p = subtitle.text_frame.paragraphs[0]

    p.text = presentation["subtitle"]

    p.font.size = Pt(18)
    p.font.color.rgb = YELLOW

    footer = slide.shapes.add_textbox(
        Inches(0.82),
        Inches(6.4),
        Inches(9),
        Inches(0.5)
    )

    p = footer.text_frame.paragraphs[0]

    p.text = (
        "AI-generated from your uploaded learning material"
    )

    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(
        203,
        213,
        225
    )

    return slide


def create_content_slide(
    prs,
    slide_data,
    slide_number
):

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_slide_background(
        slide,
        CREAM
    )

    add_top_bar(slide)

    add_title(
        slide,
        slide_data["title"],
        slide_data["subtitle"]
    )

    slide_type = slide_data["type"]

    if slide_type in {
        "concept",
        "objectives",
        "key_insight",
        "summary",
        "case_study"
    }:

        points = slide_data["points"]

        if not points:

            points = [
                slide_data["insight"]
            ] if slide_data["insight"] else [
                "Important information from the uploaded learning material."
            ]

        half = max(
            1,
            len(points) // 2
        )

        left_points = points[:half]
        right_points = points[half:]

        add_bullet_card(
            slide,
            0.65,
            1.65,
            5.9,
            4.9,
            "Core Concepts",
            left_points,
            YELLOW
        )

        add_bullet_card(
            slide,
            6.75,
            1.65,
            5.9,
            4.9,
            "Important Details",
            right_points or left_points,
            BLUE
        )

    elif slide_type == "comparison":

        add_comparison_cards(
            slide,
            slide_data["items"]
        )

    elif slide_type == "table":

        add_table_slide(
            slide,
            slide_data["items"]
        )

    elif slide_type == "process":

        add_process_nodes(
            slide,
            slide_data["items"]
        )

        if slide_data["points"]:

            add_bullet_card(
                slide,
                0.8,
                4.35,
                11.7,
                1.9,
                "Process Notes",
                slide_data["points"][:4],
                YELLOW
            )

    elif slide_type == "components":

        items = slide_data["items"][:6]

        if items:

            for index, item in enumerate(items):

                row = index // 3
                col = index % 3

                x = 0.7 + col * 4.15
                y = 1.7 + row * 2.25

                add_bullet_card(
                    slide,
                    x,
                    y,
                    3.75,
                    1.85,
                    item["name"],
                    [
                        item["description"],
                        item["example"]
                    ],
                    (
                        YELLOW
                        if index % 2 == 0
                        else BLUE
                    )
                )

    elif slide_type == "timeline":

        items = slide_data["items"][:6]

        if items:

            line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(1),
                Inches(3.15),
                Inches(11),
                Inches(0.08)
            )

            line.fill.solid()
            line.fill.fore_color.rgb = NAVY
            line.line.fill.background()

            for index, item in enumerate(items):

                x = 1 + index * 2.0

                node = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL,
                    Inches(x),
                    Inches(2.8),
                    Inches(0.75),
                    Inches(0.75)
                )

                node.fill.solid()
                node.fill.fore_color.rgb = (
                    YELLOW
                    if index % 2 == 0
                    else BLUE
                )

                node.line.fill.background()

                box = slide.shapes.add_textbox(
                    Inches(x - 0.45),
                    Inches(3.7),
                    Inches(1.7),
                    Inches(1.2)
                )

                tf = box.text_frame
                tf.word_wrap = True

                p = tf.paragraphs[0]

                p.text = item["name"]

                p.alignment = PP_ALIGN.CENTER
                p.font.bold = True
                p.font.size = Pt(11)
                p.font.color.rgb = NAVY

                if item["description"]:

                    p = tf.add_paragraph()

                    p.text = item["description"]
                    p.alignment = PP_ALIGN.CENTER
                    p.font.size = Pt(9)
                    p.font.color.rgb = GRAY

    elif slide_type == "chart":

        created = add_chart_slide(
            slide,
            slide_data["data"]
        )

        if not created:

            add_bullet_card(
                slide,
                0.7,
                1.7,
                11.8,
                4.8,
                "Key Information",
                slide_data["points"],
                YELLOW
            )

    elif slide_type == "quiz":

        questions = slide_data["questions"]

        if not questions:

            questions = [
                "What is the most important concept discussed in this material?"
            ]

        for index, question in enumerate(
            questions[:4]
        ):

            y = 1.65 + index * 1.15

            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.75),
                Inches(y),
                Inches(11.8),
                Inches(0.9)
            )

            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            card.line.color.rgb = MID_GRAY

            number = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(0.95),
                Inches(y + 0.15),
                Inches(0.55),
                Inches(0.55)
            )

            number.fill.solid()
            number.fill.fore_color.rgb = YELLOW
            number.line.fill.background()

            number_box = slide.shapes.add_textbox(
                Inches(0.95),
                Inches(y + 0.27),
                Inches(0.55),
                Inches(0.2)
            )

            p = number_box.text_frame.paragraphs[0]
            p.text = str(index + 1)
            p.alignment = PP_ALIGN.CENTER
            p.font.bold = True
            p.font.size = Pt(9)

            q_box = slide.shapes.add_textbox(
                Inches(1.7),
                Inches(y + 0.18),
                Inches(10.3),
                Inches(0.5)
            )

            p = q_box.text_frame.paragraphs[0]

            p.text = question

            p.font.size = Pt(12)
            p.font.color.rgb = NAVY

        if slide_data["answer"]:

            answer_box = slide.shapes.add_textbox(
                Inches(0.9),
                Inches(6.1),
                Inches(11.4),
                Inches(0.55)
            )

            p = answer_box.text_frame.paragraphs[0]

            p.text = (
                "Answer guidance: "
                + slide_data["answer"]
            )

            p.font.size = Pt(10)
            p.font.color.rgb = GRAY

    elif slide_type == "references":

        points = slide_data["points"]

        if not points:

            points = [
                slide_data["source"]
            ] if slide_data["source"] else [
                "References identified from the uploaded material."
            ]

        add_bullet_card(
            slide,
            0.75,
            1.65,
            11.8,
            4.9,
            "Sources & References",
            points,
            PURPLE
        )

    else:

        add_bullet_card(
            slide,
            0.75,
            1.65,
            11.8,
            4.9,
            "Important Information",
            slide_data["points"],
            YELLOW
        )

    add_footer(
        slide,
        slide_number
    )

    return slide


# ============================================================
# CREATE PPTX
# ============================================================

def create_powerpoint(
    presentation_data: Dict[str, Any],
    original_filename: str
) -> Dict[str, Any]:

    try:

        prs = Presentation()

        # 16:9 widescreen
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        title_slide = create_title_slide(
            prs,
            presentation_data,
            1
        )

        slides = presentation_data.get(
            "slides",
            []
        )

        for index, slide_data in enumerate(
            slides,
            start=2
        ):

            create_content_slide(
                prs,
                slide_data,
                index
            )

        base_name = Path(
            original_filename
        ).stem

        base_name = clean_filename(
            base_name
        )

        output_name = (
            f"STATWISE_AI_{base_name}_"
            f"{os.getpid()}.pptx"
        )

        output_path = (
            GENERATED_DIR / output_name
        )

        prs.save(
            str(output_path)
        )

        return {
            "success": True,
            "filename": output_name,
            "path": str(output_path),
            "slides": len(prs.slides),
            "concepts": len(
                presentation_data.get(
                    "slides",
                    []
                )
            ),
            "visuals": sum(
                1
                for slide in presentation_data.get(
                    "slides",
                    []
                )
                if slide.get("type") in {
                    "chart",
                    "comparison",
                    "process",
                    "timeline",
                    "components",
                    "table"
                }
            ),
        }

    except Exception as error:

        print(
            "\n========== PPT CREATION ERROR ==========\n",
            error
        )

        return {
            "success": False,
            "message": "Failed to create PowerPoint file.",
            "error": str(error)
        }


# ============================================================
# MAIN SERVICE
# ============================================================

def generate_presentation(
    pdf_text: str,
    original_filename: str,
    style: str = "Visual Learning"
) -> Dict[str, Any]:

    plan_result = generate_presentation_plan(
        pdf_text=pdf_text,
        style=style
    )

    if not plan_result.get("success"):
        return plan_result

    presentation = plan_result[
        "presentation"
    ]

    ppt_result = create_powerpoint(
        presentation_data=presentation,
        original_filename=original_filename
    )

    if not ppt_result.get("success"):
        return ppt_result

    return {
        "success": True,
        "filename": ppt_result["filename"],
        "presentation": presentation,
        "slides": ppt_result["slides"],
        "concepts": ppt_result["concepts"],
        "visuals": ppt_result["visuals"],
    }